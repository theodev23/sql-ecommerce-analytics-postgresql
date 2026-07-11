"""Convert the Online Retail Excel dataset to a PostgreSQL-friendly CSV file."""

from __future__ import annotations

import csv
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
INPUT_FILE = PROJECT_ROOT / "data" / "online_retail.xlsx"
OUTPUT_FILE = PROJECT_ROOT / "data" / "online_retail.csv"


def format_value(value: Any) -> Any:
    """Format Excel values without applying business transformations."""
    if value is None:
        return ""

    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")

    if isinstance(value, date):
        return value.isoformat()

    return value


def main() -> None:
    """Read the Excel workbook and write its active sheet as UTF-8 CSV."""
    if not INPUT_FILE.exists():
        raise FileNotFoundError(f"Source file not found: {INPUT_FILE}")

    workbook = load_workbook(
        filename=INPUT_FILE,
        read_only=True,
        data_only=True,
    )
    worksheet = workbook.active

    row_count = 0

    try:
        with OUTPUT_FILE.open(
            mode="w",
            encoding="utf-8",
            newline="",
        ) as csv_file:
            writer = csv.writer(csv_file)

            for row in worksheet.iter_rows(values_only=True):
                writer.writerow([format_value(value) for value in row])
                row_count += 1
    finally:
        workbook.close()

    print(f"CSV created: {OUTPUT_FILE}")
    print(f"Rows written including header: {row_count}")
    print(f"Data rows: {max(row_count - 1, 0)}")


if __name__ == "__main__":
    main()
